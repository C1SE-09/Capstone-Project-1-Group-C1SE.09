"""
app/services/health_exam_service.py
====================================
Dịch vụ tạo file Excel (.xlsx) mẫu sổ sức khoẻ định kỳ cho học sinh:
- Điền sẵn thông tin định danh học sinh: Họ và tên (hiển thị dạng
  "Nguyễn Văn Ánh (Nhỏ)" — tức Họ tên kèm biệt danh nếu có), Ngày sinh, Giới tính.
- KHÔNG có cột mã học sinh: khi nhập file, hệ thống đối chiếu học sinh bằng
  Họ tên + Biệt danh (ô "Họ và tên" trong file mẫu).
- Điền sẵn số liệu đã có sẵn của học sinh (Chiều cao, Cân nặng, Kết luận chung)
  nếu bé đã có bản ghi sức khoẻ trước đó, để người dùng chỉ cần sửa chỗ khác.
  BMI và phân loại thể trạng do hệ thống tự tính khi nhập file nên không cần cột riêng.
- Format cơ bản giống file mẫu giáo viên / học sinh (tiêu đề + dòng dữ liệu,
  không màu mè, không khóa sheet): người dùng tự do sửa/xóa/thêm dòng.
"""

from __future__ import annotations

import re
from datetime import date
from io import BytesIO
from typing import TYPE_CHECKING, Optional

import openpyxl
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from app.models.child import Child
    from app.models.health_examination import HealthExamination
    from app.models.user import Class_

HEADERS = [
    "Họ và tên",
    "Ngày sinh",
    "Giới tính",
    "Ngày ghi nhận",
    "Chiều cao (cm)",
    "Cân nặng (kg)",
    "Kết luận chung",
]

# Độ rộng cột cơ bản (giống file mẫu giáo viên / học sinh).
COL_WIDTHS = [38, 15, 12, 15, 16, 16, 35]


def sanitize_sheet_title(name: str) -> str:
    """Loại bỏ ký tự cấm trong tên sheet Excel và giới hạn tối đa 31 ký tự."""
    clean = re.sub(r"[\\/?*\[\]:]", "-", name).strip()
    return clean[:31] or "SoSucKhoe"


def _gender_vi(gender: object) -> str:
    g = str(gender or "").strip().lower()
    if g == "male":
        return "Nam"
    if g == "female":
        return "Nữ"
    return str(gender or "")


def display_name(full_name: str, nickname: str | None) -> str:
    """Tên hiển thị trong file mẫu: "Nguyễn Văn Ánh (Nhỏ)" (bỏ phần ngoặc nếu không có biệt danh).

    Đây cũng chính là giá trị dùng để đối chiếu học sinh khi nhập file.
    """
    name = (full_name or "").strip()
    nick = (nickname or "").strip()
    if not nick:
        return name
    return f"{name} ({nick})"


def _fill_sheet(
    ws,
    children: list[Child],
    default_date: Optional[date] = None,
    existing_exams: Optional[dict[str, HealthExamination]] = None,
) -> None:
    """Đổ tiêu đề + dòng dữ liệu vào một sheet (format cơ bản, không style màu)."""
    ws.append(HEADERS)
    exam_date_str = (default_date or date.today()).strftime("%d/%m/%Y")
    existing_exams = existing_exams or {}
    for child in children:
        dob_str = child.date_of_birth.strftime("%d/%m/%Y") if child.date_of_birth else ""
        record = existing_exams.get(str(child.child_id))
        if record is not None:
            # Đã có số liệu sức khoẻ: điền sẵn để người dùng rà soát / sửa.
            row_exam_date = record.examination_date.strftime("%d/%m/%Y")
            height: object = record.height
            weight: object = record.weight
            conclusion: object = record.general_conclusion or None
        else:
            row_exam_date = exam_date_str
            height = None  # Chiều cao (cm) — để trống cho người dùng điền
            weight = None  # Cân nặng (kg) — để trống cho người dùng điền
            conclusion = None  # Kết luận chung — để trống
        ws.append([
            display_name(child.full_name, child.nickname),
            dob_str,
            _gender_vi(child.gender),
            row_exam_date,
            height,
            weight,
            conclusion,
        ])
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def generate_health_exam_template(
    classroom: Class_,
    children: list[Child],
    default_date: Optional[date] = None,
    existing_exams: Optional[dict[str, HealthExamination]] = None,
) -> bytes:
    """Tạo file Excel (.xlsx) mẫu cho 1 lớp (1 sheet, format cơ bản)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sanitize_sheet_title(f"So_Suc_Khoe_{classroom.name}")
    _fill_sheet(ws, children, default_date, existing_exams)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def generate_health_exam_templates_all(
    classrooms: list[tuple[Class_, list[Child]]],
    default_date: Optional[date] = None,
    existing_exams: Optional[dict[str, HealthExamination]] = None,
) -> bytes:
    """Tạo file Excel mẫu cho tất cả lớp — mỗi lớp 1 sheet, format cơ bản."""
    wb = openpyxl.Workbook()
    first = True
    used_titles: set[str] = set()
    for classroom, children in classrooms:
        ws = wb.active if first else wb.create_sheet()
        first = False
        title = sanitize_sheet_title(classroom.name)
        # Tránh trùng tên sheet (Excel không cho 2 sheet cùng tên).
        base, counter = title, 2
        while title.lower() in used_titles:
            suffix = f" ({counter})"
            title = base[: 31 - len(suffix)] + suffix
            counter += 1
        used_titles.add(title.lower())
        ws.title = title
        _fill_sheet(ws, children, default_date, existing_exams)
    if not classrooms:
        ws = wb.active
        ws.title = "SoSucKhoe"
        _fill_sheet(ws, [], default_date, existing_exams)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
